"""bytes + filename -> sections. One small parser per format.

PRODUCTION: Document AI Layout Parser for scanned PDFs, tables, multi-column layouts. Add it as
another entry in PARSERS; nothing downstream changes.
"""

import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import PurePath


class UnsupportedFormat(ValueError):
    pass


@dataclass
class Section:
    heading: str
    text: str
    page: int | None = None


@dataclass
class ParsedDocument:
    format: str
    sections: list[Section]
    page_count: int | None = None
    title_hint: str = ""
    outline: list[str] = field(default_factory=list)

    def full_text(self) -> str:
        return "\n\n".join(f"{s.heading}\n{s.text}".strip() for s in self.sections)


MIME = {
    "md": "text/markdown",
    "txt": "text/plain",
    "html": "text/html",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "csv": "text/csv",
    "json": "application/json",
}
ALIASES = {"markdown": "md", "htm": "html", "text": "txt"}
RULE_FORMATS = {"csv", "json"}


def format_of(filename: str) -> str:
    ext = PurePath(filename).suffix.lower().lstrip(".")
    return ALIASES.get(ext, ext)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _finish(fmt: str, sections: list[Section], **kw) -> ParsedDocument:
    sections = [s for s in sections if s.text.strip() or s.heading.strip()]
    outline = [s.heading for s in sections if s.heading and not s.heading.startswith("Page ")]
    title = kw.pop("title_hint", "") or (outline[0] if outline else "")
    return ParsedDocument(fmt, sections, title_hint=title, outline=outline[:50], **kw)


def parse_markdown(data: bytes) -> ParsedDocument:
    sections, heading, buf = [], "", []
    for line in _decode(data).splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            sections.append(Section(heading, "\n".join(buf).strip()))
            heading, buf = m.group(2).strip(), []
        else:
            buf.append(line)
    sections.append(Section(heading, "\n".join(buf).strip()))
    return _finish("md", sections)


_TXT_HEADING = re.compile(r"^(\d+(\.\d+)*[.)]?\s+\S.{0,78}|[A-Z][A-Z0-9 ,&/-]{3,78})$")


def parse_text(data: bytes) -> ParsedDocument:
    sections, heading, buf = [], "", []
    for block in re.split(r"\n\s*\n", _decode(data)):
        block = block.strip()
        if not block:
            continue
        first, _, rest = block.partition("\n")
        if _TXT_HEADING.match(first.strip()):
            sections.append(Section(heading, "\n\n".join(buf)))
            heading, buf = first.strip(), [rest.strip()] if rest.strip() else []
        else:
            buf.append(block)
    sections.append(Section(heading, "\n\n".join(buf)))
    return _finish("txt", sections)


class _HtmlSections(HTMLParser):
    HEADINGS = {"h1", "h2", "h3", "h4"}
    BLOCKS = {"p", "li", "tr", "div", "br", "section", "article", "td", "th"}
    SKIP = {"script", "style", "noscript", "head"}

    def __init__(self) -> None:
        super().__init__()
        self.sections: list[Section] = []
        self.heading, self.buf, self.in_heading, self.skip, self.title = "", [], False, 0, ""
        self._heading_buf: list[str] = []
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif tag == "title":
            self._in_title = True
        elif tag in self.HEADINGS:
            self.sections.append(Section(self.heading, "".join(self.buf).strip()))
            self.buf, self.in_heading, self._heading_buf = [], True, []
        elif tag in self.BLOCKS:
            self.buf.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
        elif tag == "title":
            self._in_title = False
        elif tag in self.HEADINGS and self.in_heading:
            self.heading, self.in_heading = " ".join("".join(self._heading_buf).split()), False

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif self.skip:
            return
        elif self.in_heading:
            self._heading_buf.append(data)
        else:
            self.buf.append(data)

    def result(self) -> list[Section]:
        self.sections.append(Section(self.heading, "".join(self.buf).strip()))
        for s in self.sections:
            s.text = re.sub(r"\n\s*\n+", "\n\n", re.sub(r"[ \t]+", " ", s.text)).strip()
        return self.sections


def parse_html(data: bytes) -> ParsedDocument:
    p = _HtmlSections()
    p.feed(_decode(data))
    return _finish("html", p.result(), title_hint=p.title.strip())


def parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    sections = [
        Section(f"Page {i}", (page.extract_text() or "").strip(), page=i)
        for i, page in enumerate(reader.pages, start=1)
    ]
    title = (reader.metadata.title if reader.metadata else "") or ""
    if not any(s.text for s in sections):
        raise UnsupportedFormat("PDF has no extractable text (scanned?). Use Document AI OCR.")
    return _finish("pdf", sections, page_count=len(reader.pages), title_hint=title)


def parse_docx(data: bytes) -> ParsedDocument:
    import docx

    d = docx.Document(io.BytesIO(data))
    sections, heading, buf = [], "", []
    for para in d.paragraphs:
        style = (para.style.name if para.style else "") or ""
        if style.startswith("Heading") or style == "Title":
            sections.append(Section(heading, "\n".join(buf).strip()))
            heading, buf = para.text.strip(), []
        elif para.text.strip():
            buf.append(para.text.strip())
    for table in d.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        buf.append("\n".join(rows))
    sections.append(Section(heading, "\n".join(buf).strip()))
    return _finish("docx", sections, title_hint=d.core_properties.title or "")


PARSERS = {
    "md": parse_markdown,
    "txt": parse_text,
    "html": parse_html,
    "pdf": parse_pdf,
    "docx": parse_docx,
}


def parse(data: bytes, filename: str) -> ParsedDocument:
    fmt = format_of(filename)
    if fmt in RULE_FORMATS:
        raise UnsupportedFormat(
            f".{fmt} files are rule tables: import them via /api/v1/rules/import"
        )
    parser = PARSERS.get(fmt)
    if parser is None:
        raise UnsupportedFormat(f"unsupported format .{fmt}; supported: {', '.join(PARSERS)}")
    return parser(data)
