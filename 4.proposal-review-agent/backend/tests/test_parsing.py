import io

import pytest

from app.rag.chunker import chunk_sections
from app.rag.parsers import Section, UnsupportedFormat, format_of, parse
from tests.conftest import SAMPLES

DOCS = SAMPLES / "documents"


@pytest.mark.parametrize(
    "name, fmt, title, heading",
    [
        ("product-launch-policy.md", "md", "Product Launch Policy", "1. Security review"),
        ("data-privacy-standard.html", "html", "Data Privacy Standard", "3. Children"),
        (
            "marketing-claims-guidelines.txt",
            "txt",
            "MARKETING CLAIMS GUIDELINES",
            "3. Comparative claims",
        ),
        ("ai-usage-policy.docx", "docx", "AI Usage Policy", "2. Training data"),
    ],
)
def test_sample_formats(name, fmt, title, heading):
    p = parse((DOCS / name).read_bytes(), name)
    assert p.format == fmt and p.title_hint == title
    assert heading in p.outline


def test_pdf_pages_and_metadata_title():
    p = parse((DOCS / "security-baseline.pdf").read_bytes(), "security-baseline.pdf")
    assert p.page_count == 2 and p.title_hint == "Security Baseline"
    assert "AES-256" in p.sections[0].text and p.sections[1].page == 2


def test_html_drops_scripts():
    p = parse((DOCS / "data-privacy-standard.html").read_bytes(), "x.html")
    assert "ignored by the parser" not in p.full_text()
    assert "verifiable parental consent" in p.full_text()


def test_docx_tables_are_kept():
    p = parse((DOCS / "ai-usage-policy.docx").read_bytes(), "x.docx")
    assert "Support reply drafts | Yes, with disclosure" in p.full_text()


def test_generated_pdf_roundtrip():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 8, "Rule: badges must be worn.", new_x="LMARGIN", new_y="NEXT")
    p = parse(bytes(pdf.output()), "gen.pdf")
    assert "badges must be worn" in p.full_text()


def test_empty_pdf_is_rejected_with_ocr_hint():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    with pytest.raises(UnsupportedFormat, match="OCR"):
        parse(bytes(pdf.output()), "scan.pdf")


def test_unsupported_and_rule_formats():
    with pytest.raises(UnsupportedFormat, match="unsupported"):
        parse(b"x", "a.xlsx")
    with pytest.raises(UnsupportedFormat, match="rule tables"):
        parse(b"a,b", "rules.csv")
    assert format_of("A.MARKDOWN") == "md" and format_of("x.htm") == "html"


def test_latin1_text_is_decoded():
    assert "café" in parse("café policy".encode("latin-1"), "x.txt").full_text()


def test_chunker_respects_sections_and_size():
    long = "\n".join(f"Sentence number {i} about compliance." for i in range(200))
    chunks = chunk_sections(
        [Section("A", long), Section("B", "short text", page=3)], max_chars=300, overlap_chars=40
    )
    assert all(len(c.text) <= 300 + 40 for c in chunks)
    assert {c.heading for c in chunks} == {"A", "B"}
    assert chunks[-1].text == "short text" and chunks[-1].page == 3


def test_docx_bytes_from_memory():
    import docx

    d = docx.Document()
    d.add_heading("Heading One", 1)
    d.add_paragraph("Body text.")
    buf = io.BytesIO()
    d.save(buf)
    p = parse(buf.getvalue(), "m.docx")
    assert p.outline == ["Heading One"] and "Body text." in p.full_text()
